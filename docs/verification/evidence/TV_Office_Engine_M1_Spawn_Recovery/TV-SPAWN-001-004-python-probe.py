from __future__ import annotations
import hashlib, importlib.metadata, io, json, multiprocessing, os, platform, resource, sys, time
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[4]
ADDONS = str(ROOT / 'mymodules')
if ADDONS not in sys.path:
    sys.path.insert(0, ADDONS)
sys.path.insert(0, str(ROOT / 'mymodules' / 'office_core' / 'tests'))
from test_office_core import make_docx

from office_core import ResourceBudget, OfficeCoreError
from office_core.engine import inspect_inline, render_inline
from office_core.service import inspect_document
from office_verify import VerificationBudget, VerificationPolicy, verify_output
from office_verify.service import _process_rss_bytes, _stop

XLSX_MIME='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

def sample_xlsx():
    wb=Workbook(); ws=wb.active; ws.title='Sheet1'; ws['A1']='old'; ws['B1']='keep'; wb.properties.created=datetime(2020,1,1); wb.properties.modified=datetime(2020,1,1)
    b=io.BytesIO(); wb.save(b); return b.getvalue()

def budget(memory=512*1024*1024, timeout=15):
    return ResourceBudget(max_file_size=2_000_000,max_expanded_size=8_000_000,max_zip_entries=128,max_entry_size=2_000_000,max_xml_elements=100_000,max_cells=100_000,max_paragraphs=10_000,max_memory_bytes=memory,timeout_seconds=timeout)

def identity_worker(send):
    import office_core, office_core.models, office_core.service
    import office_verify, office_verify.models, office_verify.service
    send.send({
      'office_core_file': office_core.__file__,
      'office_core_service_module': office_core.service.__name__,
      'resource_budget_class_module': office_core.ResourceBudget.__module__,
      'office_verify_file': office_verify.__file__,
      'office_verify_service_module': office_verify.service.__name__,
      'verification_budget_class_module': office_verify.VerificationBudget.__module__,
      'sys_path_addons_count': sys.path.count(ADDONS),
    })
    send.close()

def core_inline_worker(send, operation, args, bgt):
    try:
        from office_core.engine import inspect_inline as inspect_fn, render_inline as render_fn
        if operation=='inspect': result=inspect_fn(*args, bgt)
        elif operation=='render': result=render_fn(*args, bgt)
        else: raise RuntimeError('invalid operation')
        send.send(('ok', result))
    except BaseException as exc:
        try: send.send(('error', type(exc).__name__, getattr(exc,'code',None), getattr(exc,'stage',None), str(exc)[:200]))
        except BaseException: pass
    finally: send.close()

def monitored_call(target,args,limit_bytes,timeout=15):
    ctx=multiprocessing.get_context('spawn'); recv,send=ctx.Pipe(duplex=False)
    proc=ctx.Process(target=target,args=(send,*args),daemon=True)
    start=time.monotonic(); proc.start(); send.close(); samples=[]; max_rss=0; limit_seen=None; payload=None; timed_out=False
    deadline=start+timeout
    while time.monotonic()<deadline:
        if proc.is_alive():
            rss=_process_rss_bytes(proc.pid)
            if rss is None:
                _stop(proc); raise RuntimeError('rss unavailable')
            samples.append((time.monotonic()-start,rss)); max_rss=max(max_rss,rss)
            if rss>limit_bytes:
                limit_seen=rss; _stop(proc); break
        if recv.poll(0.01):
            payload=recv.recv(); proc.join(0.25); break
        if not proc.is_alive():
            proc.join(0.1); break
    else:
        timed_out=True; _stop(proc)
    if proc.is_alive(): _stop(proc)
    end=time.monotonic(); alive_after=proc.is_alive(); exitcode=proc.exitcode
    recv.close()
    return {'payload':payload,'elapsed_ms':round((end-start)*1000,3),'max_sampled_rss_bytes':max_rss,'limit_seen_rss_bytes':limit_seen,'samples':len(samples),'exitcode':exitcode,'alive_after':alive_after,'timed_out':timed_out,'overshoot_bytes':max(0,(limit_seen or max_rss)-limit_bytes)}

def allocate_worker(send, max_chunks=64):
    chunks=[]
    try:
        for _ in range(max_chunks):
            chunks.append(bytearray(4*1024*1024))
            time.sleep(0.004)
        send.send('completed_without_limit')
    except BaseException as exc:
        try: send.send(type(exc).__name__)
        except BaseException: pass
    finally: send.close()

def main():
    content=sample_xlsx(); docx_bytes=make_docx()
    out={'environment':{'python':sys.version,'platform':platform.platform(),'pid':os.getpid(),'RLIMIT_AS':list(resource.getrlimit(resource.RLIMIT_AS)),'packages':{name:importlib.metadata.version(name) for name in ('openpyxl','python-docx','lxml')}},'sample_sha256':hashlib.sha256(content).hexdigest(),'docx_sample_sha256':hashlib.sha256(docx_bytes).hexdigest(),'addon_path':ADDONS}
    # Spawn package/module identity.
    ctx=multiprocessing.get_context('spawn'); recv,send=ctx.Pipe(duplex=False); p=ctx.Process(target=identity_worker,args=(send,),daemon=True); t=time.monotonic(); p.start(); send.close(); ident=recv.recv(); p.join(3); out['spawn_identity']={'result':ident,'elapsed_ms':round((time.monotonic()-t)*1000,3),'exitcode':p.exitcode}
    # Existing core public service under the actual 4 GiB budget; record exact result.
    b4=budget(4*1024*1024*1024,15)
    try:
        inspected=inspect_document(content,'sample.xlsx',XLSX_MIME,b4)
        out['existing_core_public_inspect']={'status':'PASS','format':inspected.format,'template_sha256':inspected.template_sha256}
    except OfficeCoreError as exc:
        out['existing_core_public_inspect']={'status':'FAIL','error_code':exc.code,'stage':exc.stage,'message':str(exc)}
    # Explore a candidate RSS-supervised spawn worker (not product code).
    b=budget(512*1024*1024,15)
    core=monitored_call(core_inline_worker,('inspect',(content,'sample.xlsx',XLSX_MIME),b),b.max_memory_bytes,b.timeout_seconds)
    out['rss_supervised_core_inspect']=core
    if core.get('payload') and core['payload'][0]=='ok':
        inspection=core['payload'][1]
        plan={'schema_version':'1','template_sha256':inspection.template_sha256,'writes':[{'element_id':'xlsx:Sheet1!A1','operation':'set_value','value':{'type':'string','data':'new'}}]}
        rendered=monitored_call(core_inline_worker,('render',(content,'sample.xlsx',XLSX_MIME,plan),b),b.max_memory_bytes,b.timeout_seconds)
        out['rss_supervised_core_render']={k:v for k,v in rendered.items() if k!='payload'}; out['rss_supervised_core_render']['worker_payload_summary']=(rendered.get('payload')[:5] if rendered.get('payload') and rendered['payload'][0] != 'ok' else ('ok', type(rendered['payload'][1]).__name__, getattr(rendered['payload'][1], 'format', None), getattr(rendered['payload'][1], 'candidate_sha256', None), getattr(rendered['payload'][1], 'applied_writes', None)))
        if rendered.get('payload') and rendered['payload'][0]=='ok':
            candidate=rendered['payload'][1]
            vb=VerificationBudget(max_file_size=2_000_000,max_expanded_size=8_000_000,max_zip_entries=128,max_entry_size=2_000_000,max_xml_elements=100_000,max_cells=100_000,max_paragraphs=10_000,max_mapping_bytes=1_000_000,max_memory_bytes=512*1024*1024,timeout_seconds=15)
            t=time.monotonic()
            try:
                report=verify_output(content,candidate.content,'sample.xlsx',XLSX_MIME,plan,VerificationPolicy(enabled_formats=('xlsx','docx')),vb)
                out['rss_supervised_verify']={'status':report.status,'elapsed_ms':round((time.monotonic()-t)*1000,3),'output_sha256':report.output_sha256,'checked_targets':report.checked_targets,'checked_formulas':report.checked_formulas,'rss_worker_monitor':'verify_output built-in RSS sampler'}
            except BaseException as exc:
                out['rss_supervised_verify']={'status':'FAIL','elapsed_ms':round((time.monotonic()-t)*1000,3),'error_type':type(exc).__name__,'error_code':getattr(exc,'code',None),'stage':getattr(exc,'stage',None),'message':str(exc)}
    # Repeat the synthetic inspect/render/independent-readback path for DOCX.
    docx_sha=hashlib.sha256(docx_bytes).hexdigest(); DOCX_MIME='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    di=monitored_call(core_inline_worker,('inspect',(docx_bytes,'sample.docx',DOCX_MIME),b),b.max_memory_bytes,b.timeout_seconds)
    out['docx_rss_supervised_inspect']={k:v for k,v in di.items() if k!='payload'}
    if di.get('payload') and di['payload'][0]=='ok':
        dinspection=di['payload'][1]
        dplan={'schema_version':'1','template_sha256':dinspection.template_sha256,'writes':[{'element_id':'docx:p:1','operation':'set_value','value':{'type':'string','data':'TV generated'}}]}
        dr=monitored_call(core_inline_worker,('render',(docx_bytes,'sample.docx',DOCX_MIME,dplan),b),b.max_memory_bytes,b.timeout_seconds)
        out['docx_rss_supervised_render']={k:v for k,v in dr.items() if k!='payload'}
        out['docx_rss_supervised_render']['worker_payload_summary']=(dr.get('payload')[:5] if dr.get('payload') and dr['payload'][0] != 'ok' else ('ok', type(dr['payload'][1]).__name__, getattr(dr['payload'][1], 'format', None), getattr(dr['payload'][1], 'candidate_sha256', None), getattr(dr['payload'][1], 'applied_writes', None)))
        if dr.get('payload') and dr['payload'][0]=='ok':
            dc=dr['payload'][1]; dvb=VerificationBudget(max_file_size=2_000_000,max_expanded_size=8_000_000,max_zip_entries=128,max_entry_size=2_000_000,max_xml_elements=100_000,max_cells=100_000,max_paragraphs=10_000,max_mapping_bytes=1_000_000,max_memory_bytes=512*1024*1024,timeout_seconds=15)
            dt=time.monotonic()
            try:
                dreport=verify_output(docx_bytes,dc.content,'sample.docx',DOCX_MIME,dplan,VerificationPolicy(enabled_formats=('xlsx','docx')),dvb)
                out['docx_rss_supervised_verify']={'status':dreport.status,'elapsed_ms':round((time.monotonic()-dt)*1000,3),'output_sha256':dreport.output_sha256,'checked_targets':dreport.checked_targets,'checked_formulas':dreport.checked_formulas}
            except BaseException as exc:
                out['docx_rss_supervised_verify']={'status':'FAIL','elapsed_ms':round((time.monotonic()-dt)*1000,3),'error_type':type(exc).__name__,'error_code':getattr(exc,'code',None),'stage':getattr(exc,'stage',None),'message':str(exc)}
    # Verify the production verifier fails closed when its RSS budget is deliberately below process baseline.
    low_vb=VerificationBudget(max_file_size=2_000_000,max_expanded_size=8_000_000,max_zip_entries=128,max_entry_size=2_000_000,max_xml_elements=100_000,max_cells=100_000,max_paragraphs=10_000,max_mapping_bytes=1_000_000,max_memory_bytes=1,timeout_seconds=5)
    low_plan={'schema_version':'1','template_sha256':hashlib.sha256(content).hexdigest(),'writes':[]}
    try:
        verify_output(content,content,'sample.xlsx',XLSX_MIME,low_plan,VerificationPolicy(enabled_formats=('xlsx','docx')),low_vb)
        out['verify_below_baseline_budget']={'status':'UNEXPECTED_PASS'}
    except BaseException as exc:
        out['verify_below_baseline_budget']={'status':'REJECTED','error_type':type(exc).__name__,'error_code':getattr(exc,'code',None),'stage':getattr(exc,'stage',None),'message':str(exc)}
    # Three bounded RSS-monitor trials; 96 MiB threshold; <= 256 MiB allocated by worker.
    out['rss_monitor_trials']=[]
    for i in range(3):
        trial=monitored_call(allocate_worker,(64,),96*1024*1024,5)
        trial.pop('payload',None); trial['trial']=i+1; trial['limit_bytes']=96*1024*1024
        out['rss_monitor_trials'].append(trial)
    result=json.dumps(out,ensure_ascii=False,indent=2,default=str)
    print(result)
    evidence_dir=ROOT/'docs/verification/evidence/TV_Office_Engine_M1_Spawn_Recovery'
    evidence_dir.mkdir(parents=True,exist_ok=True)
    (evidence_dir/'TV-SPAWN-001-004-python-probe.json').write_text(result+'\n',encoding='utf-8')

if __name__=='__main__': main()

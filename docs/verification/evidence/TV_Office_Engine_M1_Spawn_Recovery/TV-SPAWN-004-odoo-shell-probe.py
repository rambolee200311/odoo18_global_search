import io, json, hashlib, uuid
from datetime import datetime
from openpyxl import Workbook
request_uuid=str(uuid.uuid4())
filename='TV-SPAWN-RECOVERY-'+request_uuid+'.xlsx'
wb=Workbook(); ws=wb.active; ws.title='Sheet1'; ws['A1']='old'; ws['B1']='keep'; wb.properties.created=datetime(2020,1,1); wb.properties.modified=datetime(2020,1,1)
buf=io.BytesIO(); wb.save(buf); content=buf.getvalue(); digest=hashlib.sha256(content).hexdigest()
attachment=env['ir.attachment'].create({'name':filename,'type':'binary','raw':content,'mimetype':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','public':False})
env.cr.commit()
plan={'schema_version':'1','template_sha256':digest,'writes':[{'element_id':'xlsx:Sheet1!A1','operation':'set_value','value':{'type':'string','data':'TV generated'}}]}
evidence={'uid':env.uid,'su':env.su,'request_uuid':request_uuid,'source_attachment_id':attachment.id,'source_sha256':digest}
try:
    result=env['office.document.request']._render_request_sync(request_uuid,attachment.id,plan)
    env.invalidate_all()
    req=env['office.document.request'].search([('request_uuid','=',request_uuid)],limit=1)
    output=req.output_attachment_id if req else env['ir.attachment'].browse([])
    evidence.update({'result':result,'request_state':req.state if req else None,'request_error_code':req.error_code if req else None,'request_failed_stage':req.failed_stage if req else None,'output_attachment_id':output.id if output else None,'output_private':(not output.public) if output else None,'output_bound_to_request':bool(output and req and output.res_model==req._name and output.res_id==req.id)})
    if output:
        output.check('read'); evidence['output_read_check']='PASS'; evidence['output_sha256']=hashlib.sha256(output.raw or b'').hexdigest()
    else:
        evidence['output_read_check']='NO_OUTPUT'
finally:
    env.invalidate_all()
    reqs=env['office.document.request'].search([('request_uuid','=',request_uuid)])
    if len(reqs)<=1:
        for req in reqs:
            if req.output_attachment_id:
                req.output_attachment_id.unlink()
            req.unlink()
        sources=env['ir.attachment'].search([('name','=',filename)])
        if len(sources)<=1:
            sources.unlink()
            env.cr.commit()
            evidence['cleanup_committed']=True
            evidence['remaining_request_count']=env['office.document.request'].search_count([('request_uuid','=',request_uuid)])
            evidence['remaining_source_count']=env['ir.attachment'].search_count([('name','=',filename)])
        else:
            env.cr.rollback(); evidence['cleanup_committed']=False; evidence['cleanup_reason']='unexpected duplicate source marker'
    else:
        env.cr.rollback(); evidence['cleanup_committed']=False; evidence['cleanup_reason']='unexpected duplicate request UUID'
print(json.dumps(evidence,ensure_ascii=False,default=str))

from __future__ import annotations

import ast
import hashlib
import multiprocessing
import sys
import struct
import warnings
import zipfile
from zipfile import ZipFile
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document
from openpyxl import Workbook
from lxml import etree

from office_core import OfficeCoreError, ResourceBudget, inspect_document, render_candidate
from office_core.engine import _parser_warning_gate


ROOT = Path(__file__).resolve().parents[3]


def budget(**overrides):
    values = {
        "max_file_size": 2_000_000,
        "max_expanded_size": 8_000_000,
        "max_zip_entries": 128,
        "max_entry_size": 2_000_000,
        "max_xml_elements": 100_000,
        "max_cells": 100_000,
        "max_paragraphs": 10_000,
        # macOS Python maps a very large shared address-space region; this is a
        # local structural-test limit only, not a production memory budget.
        "max_memory_bytes": 2_000_000_000_000,
        "timeout_seconds": 15,
    }
    values.update(overrides)
    return ResourceBudget(**values)


def make_xlsx(*, formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "old"
    if formula:
        sheet["B1"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_docx():
    document = Document()
    document.add_paragraph("Existing title")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Existing cell"
    output = BytesIO()
    document.save(output)
    source = ZipFile(BytesIO(output.getvalue()), "r")
    clean = BytesIO()
    type_ns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    with zipfile.ZipFile(clean, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in source.infolist():
            name = item.filename
            if name.startswith("customXml/") or name in {"word/stylesWithEffects.xml", "docProps/thumbnail.jpeg"}:
                continue
            raw = source.read(name)
            if name == "word/_rels/document.xml.rels":
                root = etree.fromstring(raw)
                for rel in list(root):
                    if rel.get("Type", "").endswith("/customXml") or rel.get("Type", "").endswith("/stylesWithEffects"):
                        root.remove(rel)
                raw = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name == "_rels/.rels":
                root = etree.fromstring(raw)
                for rel in list(root):
                    if rel.get("Type", "").endswith("/thumbnail"):
                        root.remove(rel)
                raw = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name == "[Content_Types].xml":
                root = etree.fromstring(raw)
                for item_node in list(root):
                    part_name = item_node.get("PartName", "")
                    if part_name in {"/customXml/itemProps1.xml", "/word/stylesWithEffects.xml"}:
                        root.remove(item_node)
                    elif item_node.tag == type_ns + "Default" and item_node.get("Extension") == "jpeg":
                        root.remove(item_node)
                raw = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(name, raw)
    return clean.getvalue()


def mapping_for(inspection, element_id, value):
    return {
        "schema_version": "1",
        "template_sha256": inspection.template_sha256,
        "writes": [{"element_id": element_id, "operation": "set_value", "value": value}],
    }


def with_extra_part(content: bytes, name: str, payload: bytes) -> bytes:
    src = zipfile.ZipFile(BytesIO(content), "r")
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr(name, payload)
    return output.getvalue()


def with_xml_child(content: bytes, member_name: str, child_xml: bytes) -> bytes:
    source = zipfile.ZipFile(BytesIO(content), "r")
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in source.infolist():
            raw = source.read(item.filename)
            if item.filename == member_name:
                root = etree.fromstring(raw)
                root.append(etree.fromstring(child_xml))
                raw = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(item, raw)
    return output.getvalue()


def corrupt_zip_member(content: bytes, member_name: str) -> bytes:
    archive = ZipFile(BytesIO(content), "r")
    info = archive.getinfo(member_name)
    data = bytearray(content)
    name_len, extra_len = struct.unpack_from("<HH", data, info.header_offset + 26)
    payload_offset = info.header_offset + 30 + name_len + extra_len
    data[payload_offset] ^= 0xFF
    return bytes(data)


class OfficeCoreTests(unittest.TestCase):
    def test_xlsx_inspection_identifies_cells_and_formulas_as_read_only(self):
        result = inspect_document(
            make_xlsx(formula=True), "source.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
        )
        items = {item.element_id: item for item in result.elements}
        self.assertIn("xlsx:Sheet1!A1", items)
        self.assertEqual(items["xlsx:Sheet1!A1"].value, "old")
        self.assertFalse(items["xlsx:Sheet1!B1"].writable)
        self.assertEqual(items["xlsx:Sheet1!B1"].kind, "formula")

    def test_docx_inspection_and_candidate_render_for_plain_paragraph_and_table(self):
        source = make_docx()
        source_hash = hashlib.sha256(source).hexdigest()
        inspected = inspect_document(
            source, "source.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", budget(),
        )
        ids = [item.element_id for item in inspected.elements]
        self.assertEqual(ids, ["docx:p:1", "docx:table:1:row:1:cell:1:p:1"])
        plan = mapping_for(inspected, ids[0], {"type": "string", "data": "Changed"})
        candidate = render_candidate(
            source, "source.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", plan, budget(),
        )
        self.assertEqual(candidate.content[:2], b"PK")
        self.assertEqual(candidate.delivery_status, "candidate_only_not_verified_or_deliverable")
        reopened = Document(BytesIO(candidate.content))
        self.assertEqual(reopened.paragraphs[0].text, "Changed")
        self.assertEqual(hashlib.sha256(source).hexdigest(), source_hash)

    def test_xlsx_candidate_render_preserves_source_and_uses_new_bytes(self):
        source = make_xlsx()
        inspected = inspect_document(
            source, "source.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
        )
        plan = mapping_for(inspected, "xlsx:Sheet1!A1", {"type": "string", "data": "changed"})
        candidate = render_candidate(
            source, "source.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", plan, budget(),
        )
        self.assertEqual(load_cell(candidate.content), "changed")
        self.assertEqual(load_cell(source), "old")
        self.assertEqual(candidate.applied_writes, 1)

    def test_formula_workbook_candidate_is_disabled_to_avoid_cache_loss(self):
        source = make_xlsx(formula=True)
        inspected = inspect_document(
            source, "source.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
        )
        plan = mapping_for(inspected, "xlsx:Sheet1!A1", {"type": "string", "data": "changed"})
        with self.assertRaises(OfficeCoreError) as caught:
            render_candidate(
                source, "source.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", plan, budget(),
            )
        self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")

    def test_mapping_rejects_unknown_missing_null_duplicate_and_formula_targets(self):
        source = make_xlsx(formula=True)
        inspected = inspect_document(
            source, "source.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
        )
        cases = [
            {"schema_version": "1", "template_sha256": inspected.template_sha256, "writes": [], "extra": True},
            {"schema_version": "1", "template_sha256": inspected.template_sha256, "writes": [{"element_id": "missing", "operation": "set_value", "value": {"type": "string", "data": "x"}}]},
            mapping_for(inspected, "xlsx:Sheet1!A1", {"type": "string", "data": None}),
            {"schema_version": "1", "template_sha256": inspected.template_sha256, "writes": [
                {"element_id": "xlsx:Sheet1!A1", "operation": "set_value", "value": {"type": "string", "data": "x"}},
                {"element_id": "xlsx:Sheet1!A1", "operation": "set_value", "value": {"type": "string", "data": "y"}},
            ]},
            mapping_for(inspected, "xlsx:Sheet1!B1", {"type": "string", "data": "overwrite formula"}),
            mapping_for(inspected, "xlsx:Sheet1!A1", {"type": "number", "data": True}),
            {"schema_version": "2", "template_sha256": inspected.template_sha256, "writes": []},
            {"schema_version": "1", "template_sha256": "0" * 64, "writes": []},
            mapping_for(inspected, "xlsx:Sheet1!A1", {"type": "string", "data": "x", "extra": 1}),
            {"schema_version": "1", "template_sha256": inspected.template_sha256, "writes": [
                {"element_id": "xlsx:Sheet1!A1", "operation": "delete", "value": {"type": "string", "data": "x"}},
            ]},
        ]
        for plan in cases:
            with self.subTest(plan=plan), self.assertRaises(OfficeCoreError) as caught:
                render_candidate(
                    source, "source.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", plan, budget(),
                )
            self.assertEqual(caught.exception.code, "OFFICE_MAPPING_INVALID")

    def test_mapping_distinguishes_empty_string_zero_and_invalid_null(self):
        source = make_xlsx()
        inspected = inspect_document(
            source, "source.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
        )
        for value, expected in (({"type": "string", "data": ""}, ""), ({"type": "number", "data": 0}, 0)):
            plan = mapping_for(inspected, "xlsx:Sheet1!A1", value)
            result = render_candidate(
                source, "source.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", plan, budget(),
            )
            self.assertEqual(load_cell(result.content), expected)

    def test_container_rejects_empty_mismatch_truncated_and_unknown_parts(self):
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        invalid = [(b"", "empty.xlsx", mime), (b"not a workbook", "fake.xlsx", mime), (make_xlsx(), "fake.docx", mime)]
        for content, filename, declared in invalid:
            with self.subTest(filename=filename), self.assertRaises(OfficeCoreError) as caught:
                inspect_document(content, filename, declared, budget())
            self.assertEqual(caught.exception.code, "OFFICE_INVALID_OR_MISMATCHED")
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(b"PK\x03\x04truncated", "broken.xlsx", mime, budget())
        self.assertEqual(caught.exception.code, "OFFICE_INVALID_OR_MISMATCHED")
        hostile = with_extra_part(make_xlsx(), "xl/vbaProject.bin", b"MZ fake active object")
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(hostile, "active.xlsx", mime, budget())
        self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")
        corrupt = corrupt_zip_member(make_xlsx(), "xl/workbook.xml")
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(corrupt, "corrupt.xlsx", mime, budget())
        self.assertEqual(caught.exception.code, "OFFICE_INVALID_OR_MISMATCHED")
        encrypted_container = bytes.fromhex("D0CF11E0A1B11AE1") + bytes(64)
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(encrypted_container, "encrypted.xlsx", mime, budget())
        self.assertEqual(caught.exception.code, "OFFICE_INVALID_OR_MISMATCHED")

    def test_external_relationship_is_rejected(self):
        source = make_xlsx()
        archive = ZipFile(BytesIO(source), "r")
        output = BytesIO()
        rel_name = "xl/_rels/workbook.xml.rels"
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in archive.infolist():
                raw = archive.read(item.filename)
                if item.filename == rel_name:
                    root = etree.fromstring(raw)
                    rel = etree.SubElement(root, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
                    rel.set("Id", "external")
                    rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink")
                    rel.set("Target", "https://example.invalid/")
                    rel.set("TargetMode", "External")
                    raw = etree.tostring(root)
                dst.writestr(item, raw)
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(
                output.getvalue(), "external.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
            )
        self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")

    def test_disabled_legacy_and_macro_formats_are_refused(self):
        # Valid CFB signature is sufficient for format recognition; content reading stays disabled.
        cfb = bytes.fromhex("D0CF11E0A1B11AE1") + bytes(64)
        for filename, mime in (("legacy.xls", "application/vnd.ms-excel"), ("legacy.doc", "application/msword")):
            with self.subTest(filename=filename), self.assertRaises(OfficeCoreError) as caught:
                inspect_document(cfb, filename, mime, budget())
            self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")
        with self.assertRaises(OfficeCoreError):
            inspect_document(make_xlsx(), "macro.xlsm", "application/vnd.ms-excel.sheet.macroEnabled.12", budget())

    def test_hostile_xml_doctype_is_rejected_without_entity_resolution(self):
        source = make_xlsx()
        archive = zipfile.ZipFile(BytesIO(source), "r")
        payload = b'<!DOCTYPE x [<!ENTITY boom "boom">]><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&boom;</t></is></c></row></sheetData></worksheet>'
        output = BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in archive.infolist():
                dst.writestr(item, payload if item.filename == "xl/worksheets/sheet1.xml" else archive.read(item.filename))
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(
                output.getvalue(), "hostile.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
            )
        self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")

    def test_declared_file_expanded_entry_and_count_budgets(self):
        source = make_xlsx()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        for limited in (
            budget(max_file_size=10),
            budget(max_expanded_size=10),
            budget(max_zip_entries=1),
            budget(max_entry_size=10),
            budget(max_xml_elements=1),
        ):
            with self.subTest(limited=limited), self.assertRaises(OfficeCoreError) as caught:
                inspect_document(source, "source.xlsx", mime, limited)
            self.assertEqual(caught.exception.code, "OFFICE_RESOURCE_LIMIT")

    def test_invalid_policy_and_timeout_are_stable_and_cleanup_worker(self):
        with self.assertRaises(ValueError):
            budget(max_expanded_size=0).validate()
        for timeout in (float("nan"), float("inf"), 10**10000):
            with self.subTest(timeout_type=type(timeout).__name__), self.assertRaises(ValueError):
                budget(timeout_seconds=timeout).validate()
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(
                make_xlsx(), "source.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                budget(timeout_seconds=0.001),
            )
        self.assertEqual(caught.exception.code, "OFFICE_RESOURCE_LIMIT")
        self.assertEqual(multiprocessing.active_children(), [])
        if sys.platform == "darwin":
            with self.assertRaises(OfficeCoreError) as memory_caught:
                inspect_document(
                    make_xlsx(), "source.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    budget(max_memory_bytes=4_000_000_000),
                )
            self.assertEqual(memory_caught.exception.code, "OFFICE_RESOURCE_LIMIT")

    def test_synthetic_xlsx_extension_structure_fails_closed(self):
        source = with_xml_child(
            make_xlsx(),
            "xl/worksheets/sheet1.xml",
            b'<extLst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
        )
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(
                source, "synthetic-extension.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
            )
        self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")

    def test_docx_textbox_is_rejected_as_unverified_structure(self):
        source = make_docx()
        archive = ZipFile(BytesIO(source), "r")
        output = BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in archive.infolist():
                raw = archive.read(item.filename)
                if item.filename == "word/document.xml":
                    root = etree.fromstring(raw)
                    root.append(etree.fromstring(
                        b'<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:txbxContent/></w:r></w:p>'
                    ))
                    raw = etree.tostring(root)
                dst.writestr(item, raw)
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(
                output.getvalue(), "textbox.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document", budget(),
            )
        self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")

    def test_warning_classifier_fails_closed(self):
        item = warnings.WarningMessage(UserWarning("Conditional Formatting extension is unsupported and will be removed"), UserWarning, "x.py", 1)
        with self.assertRaises(OfficeCoreError) as caught:
            _parser_warning_gate([item], "warning-test")
        self.assertEqual(caught.exception.code, "OFFICE_UNSUPPORTED_FEATURE")

    def test_public_errors_are_stable_and_sanitized(self):
        with self.assertRaises(OfficeCoreError) as caught:
            inspect_document(
                b"not a document", "private-template.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", budget(),
            )
        error = caught.exception.as_dict()
        self.assertEqual(set(error), {"error_code", "failed_stage", "message"})
        self.assertEqual(error["error_code"], "OFFICE_INVALID_OR_MISMATCHED")
        self.assertNotIn("Traceback", str(error))
        self.assertNotIn(str(ROOT), str(error))

    def test_package_boundary_has_no_forbidden_imports_or_ignore_warning(self):
        package_root = Path(__file__).resolve().parents[1]
        forbidden = {"odoo", "llm", "requests", "httpx", "urllib", "subprocess"}
        for path in package_root.glob("*.py"):
            source = path.read_text(encoding="utf-8")
            tree = ast.parse(source)
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    self.assertFalse(any(alias.name.split(".")[0] in forbidden for alias in node.names), path.name)
                elif isinstance(node, ast.ImportFrom):
                    self.assertNotIn((node.module or "").split(".")[0], forbidden, path.name)
            self.assertNotIn('simplefilter("ignore")', source)
            self.assertNotIn("simplefilter('ignore')", source)


def load_cell(content: bytes):
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), data_only=False)
    cell = workbook.active["A1"]
    result = "" if cell.value is None and cell.data_type == "inlineStr" else cell.value
    workbook.close()
    return result


if __name__ == "__main__":
    unittest.main()

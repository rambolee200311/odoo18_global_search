from io import BytesIO
from uuid import uuid4

import psycopg2
from openpyxl import Workbook

from odoo.exceptions import AccessError, ValidationError
from odoo.tests import TransactionCase

from ..models.office_document_request import (
    _internal_mutation,
    _resource_budget,
    _verification_budget,
    inspect_document,
    render_candidate,
    verify_output,
    VerificationPolicy,
)


class TestOfficeDocumentRequest(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.group_user = cls.env.ref("base.group_user")
        cls.user_a = cls.env["res.users"].create(
            {
                "name": "Office Request Owner",
                "login": "office-owner-" + uuid4().hex,
                "email": "office-owner-" + uuid4().hex + "@example.invalid",
                "groups_id": [(6, 0, [cls.group_user.id])],
            }
        )
        cls.user_b = cls.env["res.users"].create(
            {
                "name": "Office Unrelated User",
                "login": "office-other-" + uuid4().hex,
                "email": "office-other-" + uuid4().hex + "@example.invalid",
                "groups_id": [(6, 0, [cls.group_user.id])],
            }
        )

    def _request(self, user=None, request_uuid=None):
        model = self.env["office.document.request"].with_user(user or self.user_a)
        with _internal_mutation():
            return model.create(
                {
                    "request_uuid": request_uuid or str(uuid4()),
                    "started_at": "2026-10-21 00:00:00",
                }
            )

    def test_minimal_fields_and_owner_is_forced(self):
        request = self._request()
        expected = {
            "request_uuid",
            "owner_id",
            "thread_id",
            "input_message_id",
            "source_attachment_id",
            "output_attachment_id",
            "state",
            "error_code",
            "failed_stage",
            "input_sha256",
            "output_sha256",
            "started_at",
            "finished_at",
        }
        self.assertTrue(expected.issubset(request._fields))
        self.assertFalse({"queue_job_id", "retry_count", "progress_percent", "mapping_plan"} & set(request._fields))
        self.assertEqual(request.owner_id, self.user_a)
        self.assertEqual(request.state, "accepted")

    def test_spawn_core_and_independent_verify_golden_path(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet["A1"] = "客户名称"
        sheet["B1"] = "待填"
        sheet["C1"] = "=1+1"
        stream = BytesIO()
        workbook.save(stream)
        source = stream.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        inspection = inspect_document(source, "spawn-smoke.xlsx", mime, _resource_budget())
        self.assertEqual(inspection.format, "xlsx")
        self.assertIn("xlsx:Sheet1!A1", {element.element_id for element in inspection.elements})

        plan = {
            "schema_version": "1",
            "template_sha256": inspection.template_sha256,
            "writes": [
                {
                    "element_id": "xlsx:Sheet1!B1",
                    "operation": "set_value",
                    "value": {"type": "string", "data": "测试客户"},
                }
            ],
        }
        candidate = render_candidate(source, "spawn-smoke.xlsx", mime, plan, _resource_budget())
        report = verify_output(
            source,
            candidate.content,
            "spawn-smoke.xlsx",
            mime,
            plan,
            VerificationPolicy(enabled_formats=("xlsx", "docx")),
            _verification_budget(),
        )
        self.assertEqual(report.status, "PASS")
        self.assertEqual(report.checked_formulas, 1)
        self.assertFalse(report.cached_value_changed)
        self.assertEqual(report.formula_cache_status[0]["source_cache"], "missing")
        self.assertEqual(report.formula_cache_status[0]["output_cache"], "missing")
        # Verify using the independent reader contract; this final value assertion
        # uses openpyxl only as test code, never as production verification.
        from openpyxl import load_workbook

        rendered = load_workbook(BytesIO(candidate.content), data_only=False)
        self.assertEqual(rendered["Sheet1"]["B1"].value, "测试客户")
        self.assertEqual(rendered["Sheet1"]["C1"].value, "=1+1")
        rendered.close()

    def test_uuid_database_constraint_rejects_duplicate(self):
        request_uuid = str(uuid4())
        self._request(request_uuid=request_uuid)
        with self.assertRaises(psycopg2.IntegrityError), self.cr.savepoint():
            self._request(request_uuid=request_uuid)

    def test_owner_rule_hides_other_users_request_and_write(self):
        request = self._request()
        other_model = self.env["office.document.request"].with_user(self.user_b)
        self.assertFalse(other_model.search([("id", "=", request.id)]))
        other_request = other_model.browse(request.id)
        with self.assertRaises(AccessError):
            other_request.check_access("read")
        with self.assertRaises(AccessError):
            other_request.write({"state": "failed"})

    def test_direct_create_and_mutation_are_rejected(self):
        with self.assertRaises(AccessError):
            self.env["office.document.request"].with_user(self.user_a).create(
                {"request_uuid": str(uuid4()), "started_at": "2026-10-21 00:00:00"}
            )
        request = self._request()
        with self.assertRaises(AccessError):
            request.with_user(self.user_a).write({"state": "validating"})

    def test_state_transition_is_guarded(self):
        request = self._request()
        with _internal_mutation():
            request.write({"state": "validating"})
        self.assertEqual(request.state, "validating")
        with self.assertRaises(ValidationError), _internal_mutation():
            request.write({"state": "generated"})

    def test_output_attachment_is_private_immutable_and_owner_scoped(self):
        request = self._request()
        with _internal_mutation():
            request.write({"state": "validating"})
            request.write({"state": "processing"})
        attachment = self.env["ir.attachment"].create(
            {
                "name": "private-office-result.xlsx",
                "type": "binary",
                "raw": b"synthetic office result",
                "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "public": False,
                "res_model": request._name,
                "res_id": request.id,
            }
        )
        with _internal_mutation():
            request.write({"output_attachment_id": attachment.id, "state": "generated"})
        self.assertFalse(attachment.public)
        attachment.with_user(self.user_a).check("read")
        with self.assertRaises(AccessError):
            attachment.with_user(self.user_a).write({"public": True})
        with self.assertRaises(AccessError):
            attachment.with_user(self.user_a).unlink()
        with self.assertRaises(AccessError):
            self.env["ir.attachment"].with_user(self.user_b).browse(attachment.id).check("read")
        with self.assertRaises(AccessError):
            self.env["ir.attachment"].with_user(self.user_b).browse(attachment.id).read(["raw"])

